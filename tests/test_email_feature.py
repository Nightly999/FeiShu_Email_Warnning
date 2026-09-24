from __future__ import annotations

import asyncio
import base64
import io
import os
import poplib
import zipfile
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from unittest.mock import AsyncMock

import pytest
from lark_oapi.ws.const import HEADER_SEQ, HEADER_SUM, HEADER_TYPE
from lark_oapi.ws.pb.pbbp2_pb2 import Frame
from openpyxl import Workbook

from app.email_pop3 import (
    EmailAuthenticationError,
    _analyze_attachment,
    count_pop3_messages,
    parse_message,
)
from app.email_security import decrypt_email_password, encrypt_email_password
from app.email_service import (
    EmailAnalysis,
    EmailPlan,
    _fallback_plan,
    _likely_needs_reply,
    _run_cancellable_analysis,
    _select_messages,
    _validated_plan,
    analyze_pending,
    build_email_login_card,
    build_email_report,
    build_email_report_card,
    cancel_email_analysis,
    is_stop_email_analysis_command,
    plan_email_request,
    sync_analyze_report,
    handle_email_command,
)
from app.feishu import TenantApp
from app.feishu_ws import (
    CardCallbackWsClient,
    dispatch_event,
    extract_email_bind_card_action,
    extract_email_page_card_action,
    process_email_page_card,
    process_email_bind_card,
    recover_processing_cards,
)
from app.logging_security import redact_sensitive_text
from app.settings import get_settings


def test_pop3_accepts_long_html_lines() -> None:
    assert poplib._MAXLINE == 10 * 1024 * 1024


def test_pop3_mailbox_count_uses_stat_and_closes_connection(monkeypatch) -> None:
    class Client:
        closed = False

        def stat(self):
            return 237, 1024

        def quit(self):
            self.closed = True

    client = Client()
    monkeypatch.setattr("app.email_pop3._login", lambda *args: client)
    assert count_pop3_messages(host="pop.example.com", port=995, username="user", password="secret", timeout=10) == 237
    assert client.closed


def test_pop3_fetches_only_selected_uidl(monkeypatch) -> None:
    class Client:
        retrieved = []
        closed = False

        def uidl(self):
            return b"+OK", [b"1 other", b"2 wanted"], 0

        def retr(self, number):
            self.retrieved.append(number)
            return b"+OK", [b"Subject: selected", b"", b"body"], 0

        def quit(self):
            self.closed = True

    client = Client()
    monkeypatch.setattr("app.email_pop3._login", lambda *_: client)
    from app.email_pop3 import fetch_message_by_uidl

    result = fetch_message_by_uidl(
        host="pop.example.com", port=995, username="user", password="secret",
        timeout=10, uidl="wanted", max_body_chars=1000,
    )
    assert result.subject == "selected"
    assert client.retrieved == [2] and client.closed


def test_email_password_is_authenticated_encrypted() -> None:
    previous = os.environ.get("EMAIL_CREDENTIAL_KEY")
    os.environ["EMAIL_CREDENTIAL_KEY"] = base64.urlsafe_b64encode(b"k" * 32).decode()
    get_settings.cache_clear()
    try:
        ciphertext = encrypt_email_password("real-password")
        assert "real-password" not in ciphertext
        assert decrypt_email_password(ciphertext) == "real-password"
        damaged = ciphertext[:-2] + ("AA" if ciphertext[-2:] != "AA" else "BB")
        try:
            decrypt_email_password(damaged)
        except Exception:
            pass
        else:
            raise AssertionError("tampered credential must not decrypt")
    finally:
        if previous is None:
            os.environ.pop("EMAIL_CREDENTIAL_KEY", None)
        else:
            os.environ["EMAIL_CREDENTIAL_KEY"] = previous
        get_settings.cache_clear()


def test_parse_email_keeps_body_and_attachment_metadata_only() -> None:
    raw = b"""From: Sender <sender@example.com>\r
To: User <user@example.com>\r
Cc: Other <other@example.com>\r
Subject: =?utf-8?b?5rWL6K+V6YKu5Lu2?=\r
Date: Tue, 15 Sep 2026 08:00:00 +0800\r
Message-ID: <mail-1@example.com>\r
MIME-Version: 1.0\r
Content-Type: multipart/mixed; boundary=x\r
\r
--x\r
Content-Type: text/plain; charset=utf-8\r
\r
Please review before Friday.\r
--x\r
Content-Type: application/pdf\r
Content-Disposition: attachment; filename=report.pdf\r
Content-Transfer-Encoding: base64\r
\r
UERGREFUQQ==\r
--x--\r
"""

    parsed = parse_message("uid-1", raw, 50_000)

    assert parsed.subject == "测试邮件"
    assert parsed.to[0]["address"] == "user@example.com"
    assert parsed.cc[0]["address"] == "other@example.com"
    assert "Please review" in parsed.text_body
    assert parsed.attachments == [
        {"name": "report.pdf", "mime_type": "application/pdf", "size": 7}
    ]


def test_excel_and_word_attachments_are_analyzed() -> None:
    excel = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单"
    sheet.append(["客户", "金额"])
    sheet.append(["华东客户", 1200])
    workbook.save(excel)

    word = io.BytesIO()
    with zipfile.ZipFile(word, "w") as archive:
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="x"><w:body><w:p><w:r><w:t>交付日期为本周五</w:t>'
            "</w:r></w:p></w:body></w:document>",
        )

    assert "华东客户" in _analyze_attachment("订单.xlsx", excel.getvalue())
    assert "交付日期为本周五" in _analyze_attachment("说明.docx", word.getvalue())


def test_excel_analysis_ignores_styled_empty_rows_and_summarizes_answers() -> None:
    excel = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["AI培训反馈收集表"])
    sheet.append(["姓名", "部门", "满意度", "是否继续使用"])
    for index in range(18):
        sheet.append([f"员工{index}", "计划一部" if index < 12 else "业务部", "高", "是"])
    sheet.cell(row=47, column=1).number_format = "0"
    workbook.save(excel)

    analysis = _analyze_attachment("反馈.xlsx", excel.getvalue()) or ""

    assert "18 条有效数据，4 个字段" in analysis
    assert "计划一部 12" in analysis
    assert "是否继续使用=是 18" in analysis
    assert "47 行" not in analysis


def test_email_fallback_plan_applies_scope_and_bounds() -> None:
    settings = get_settings()
    assert EmailPlan(action="query").limit == 999
    plan = _fallback_plan("查看最近3天抄送给我的重要邮件", settings)
    assert plan.lookback_hours == 72
    assert plan.scope == "cc"
    assert plan.important_only is True

    retention = _fallback_plan("邮件分析保留999天", settings)
    assert retention.action == "retention"
    assert retention.retention_days == settings.email_max_retention_days
    assert _fallback_plan("重新登录邮箱", settings).action == "rebind"
    assert _fallback_plan("登录邮箱", settings).action == "rebind"
    assert _fallback_plan("邮箱登录", settings).action == "rebind"

    hourly = _fallback_plan("分析最近12小时的前5封未处理邮件", settings)
    assert hourly.lookback_hours == 12
    assert hourly.limit == 5
    assert hourly.only_unprocessed is True

    reply_needed = _fallback_plan("告诉我哪些邮件我没有回复", settings)
    assert reply_needed.reply_needed_only is True

    history = _fallback_plan("帮我分析历史邮件前80份", settings)
    assert history.lookback_hours == 24 * 365 * 10
    assert history.limit == 80
    assert _fallback_plan("有多少封重要邮件", settings).action == "query"

    grouped_text = "请把这100封邮件按照直接写给我的和抄送给我的分成两类"
    grouped = _fallback_plan(grouped_text, settings)
    assert grouped.action == "query"
    assert grouped.limit == 100
    assert grouped.scope == "to_or_cc"
    assert grouped.group_by_recipient is True
    assert _validated_plan(
        EmailPlan(action="unrelated"), settings, grouped_text
    ).group_by_recipient is True


def test_stop_email_analysis_cancels_only_the_current_user(monkeypatch) -> None:
    event = {"tenant_key": "tenant", "app_id": "app", "open_id": "user"}
    started = asyncio.Event()

    async def slow_analysis(*_args):
        started.set()
        await asyncio.sleep(30)

    async def scenario():
        monkeypatch.setattr("app.email_service.sync_analyze_report", slow_analysis)
        runner = asyncio.create_task(
            _run_cancellable_analysis(event, {}, EmailPlan(action="query"))
        )
        await started.wait()
        assert cancel_email_analysis(event) is True
        result = await runner
        assert result.answer == "邮件分析已停止。"

    asyncio.run(scenario())


@pytest.mark.parametrize(
    "text",
    [
        "取消指令",
        "取消分析",
        "取消分类",
        "停止整理",
        "终止邮件处理",
        "暂停当前任务",
        "结束这次操作",
        "取消",
        "算了",
        "不用了",
        "别再分析了",
    ],
)
def test_stop_email_analysis_accepts_natural_phrasing(text: str) -> None:
    assert is_stop_email_analysis_command(text) is True


@pytest.mark.parametrize(
    "text",
    [
        "停止定时任务",
        "取消每天推送",
        "取消邮箱绑定",
        "不要停止分析",
        "分析我的邮件",
    ],
)
def test_stop_email_analysis_keeps_other_commands(text: str) -> None:
    assert is_stop_email_analysis_command(text) is False


def test_email_analysis_timeout_returns_retry_guidance(monkeypatch) -> None:
    event = {"tenant_key": "tenant", "app_id": "app", "open_id": "user"}

    async def slow_analysis(*_args):
        await asyncio.sleep(30)

    monkeypatch.setattr("app.email_service.sync_analyze_report", slow_analysis)
    monkeypatch.setattr(
        "app.email_service.get_settings",
        lambda: type("Settings", (), {"email_analysis_timeout_seconds": 0.01})(),
    )

    result = asyncio.run(
        _run_cancellable_analysis(event, {}, EmailPlan(action="query"))
    )

    assert "已自动结束" in result.answer
    assert "最近20封有附件的邮件" in result.answer


def test_text_processing_exception_replaces_progress_card_with_input_guide(monkeypatch) -> None:
    app = TenantApp("tenant", "app", "secret", None, None, "bot", None)
    monkeypatch.setattr(
        "app.feishu_ws.get_active_session_id", AsyncMock(return_value="session")
    )
    monkeypatch.setattr("app.feishu_ws.hydrate_reply_context", AsyncMock())
    monkeypatch.setattr(
        "app.feishu_ws.reply_card", AsyncMock(return_value="progress-message")
    )
    monkeypatch.setattr("app.feishu_ws.record_event_progress", AsyncMock())
    monkeypatch.setattr(
        "app.feishu_ws.handle_builtin_text_command",
        AsyncMock(side_effect=RuntimeError("boom")),
    )
    update = AsyncMock(return_value=True)
    monkeypatch.setattr("app.feishu_ws.update_card", update)
    event = {
        "tenant_key": "tenant",
        "app_id": "app",
        "bot_code": "bot",
        "open_id": "user",
        "chat_id": "chat",
        "message_id": "message",
        "message_type": "text",
        "text": "帮我处理邮件",
    }

    with pytest.raises(RuntimeError, match="boom"):
        asyncio.run(dispatch_event(app, event))

    guide_card = update.await_args.args[2]
    assert update.await_args.args[:2] == (app, "progress-message")
    assert guide_card["header"]["title"]["content"] == "请完善指令"
    assert guide_card["header"]["template"] == "orange"
    content = guide_card["body"]["elements"][0]["content"]
    assert "动作 + 邮件范围 + 时间或数量 + 期望结果" in content
    assert "每天上午9点分析未处理邮件并推送给我" in content
    assert "失败" not in content
    assert "错误" not in content


def test_restart_replaces_stale_processing_card(monkeypatch) -> None:
    app = TenantApp("tenant", "app", "secret", None, None, "bot", None)
    monkeypatch.setattr(
        "app.feishu_ws.list_processing_events",
        AsyncMock(
            return_value=[
                {
                    "message_id": "request-message",
                    "progress_message_id": "progress-message",
                    "request_text": "整理有附件的邮件",
                }
            ]
        ),
    )
    update = AsyncMock(return_value=False)
    reply = AsyncMock()
    finish = AsyncMock()
    monkeypatch.setattr("app.feishu_ws.update_card", update)
    monkeypatch.setattr("app.feishu_ws.reply_card", reply)
    monkeypatch.setattr("app.feishu_ws.finish_event", finish)

    asyncio.run(recover_processing_cards(app))

    card = update.await_args.args[2]
    assert update.await_args.args[:2] == (app, "progress-message")
    assert card["header"]["title"]["content"] == "请重新发送"
    assert "服务重启而中断" in card["body"]["elements"][0]["content"]
    reply.assert_awaited_once_with(app, "request-message", card)
    finish.assert_awaited_once_with(
        tenant_key="tenant",
        app_id="app",
        message_id="request-message",
        error="service restarted",
    )


def test_explicit_email_login_skips_model_planning(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    model = AsyncMock()
    monkeypatch.setattr("app.email_service.invoke_chat_with_fallback", model)

    plan = asyncio.run(plan_email_request("登录邮箱"))

    assert plan.action == "rebind"
    model.assert_not_awaited()


@pytest.mark.parametrize(
    ("question", "kind", "hours"),
    [
        ("邮箱里一共有多少封邮件", "all", 48),
        ("邮箱里有多少邮件", "all", 48),
        ("最近48小时有多少封邮件", "recent", 48),
        ("这里一共有多少封未读邮件", "unread", 48),
    ],
)
def test_count_questions_skip_model_planning(monkeypatch, question, kind, hours) -> None:
    model = AsyncMock()
    monkeypatch.setattr("app.email_service.invoke_chat_with_fallback", model)
    plan = asyncio.run(plan_email_request(question))
    assert plan.action == "count"
    assert plan.count_kind == kind
    assert plan.lookback_hours == hours
    model.assert_not_awaited()


def test_count_commands_do_not_analyze_or_truncate(monkeypatch) -> None:
    account = {
        "id": 12,
        "email_address": "user@example.com",
        "pop3_host": "pop.example.com",
        "pop3_port": 995,
        "password_ciphertext": "encrypted",
    }
    monkeypatch.setattr("app.email_service.get_email_account", AsyncMock(return_value=account))
    monkeypatch.setattr("app.email_service.decrypt_email_password", lambda _: "password")

    def mailbox_count(**_) -> int:
        return 237

    monkeypatch.setattr("app.email_service.count_pop3_messages", mailbox_count)
    recent_count = AsyncMock(return_value=153)
    monkeypatch.setattr("app.email_service.count_recent_messages", recent_count)
    sync = AsyncMock(return_value=0)
    monkeypatch.setattr("app.email_service.sync_account", sync)
    analyze = AsyncMock()
    monkeypatch.setattr("app.email_service.analyze_pending", analyze)
    event = {"tenant_key": "tenant", "app_id": "app", "open_id": "user"}

    total = asyncio.run(handle_email_command({**event, "text": "邮箱里一共有多少封邮件"}))
    assert "237" in total.answer and "收件箱" in total.answer
    assert total.card is None
    sync.assert_not_awaited()

    recent = asyncio.run(handle_email_command({**event, "text": "最近48小时有多少封邮件"}))
    assert "153" in recent.answer and "已同步" in recent.answer
    sync.assert_awaited_once()
    recent_count.assert_awaited_once_with(12, 48)
    analyze.assert_not_awaited()

    unread = asyncio.run(handle_email_command({**event, "text": "有多少封未读邮件"}))
    assert "无法" in unread.answer and "未读" in unread.answer and "POP3" in unread.answer
    assert "153" not in unread.answer
    analyze.assert_not_awaited()


def test_bound_email_status_shows_full_address_only_in_private_chat(monkeypatch) -> None:
    account = {
        "email_address": "shxm18@example.com",
        "retention_days": 7,
        "last_sync_at": "2026-09-22 09:00:00",
    }
    lookup = AsyncMock(return_value=account)
    monkeypatch.setattr("app.email_service.get_email_account", lookup)
    monkeypatch.setattr(
        "app.email_service.plan_email_request",
        AsyncMock(return_value=EmailPlan(action="status")),
    )
    event = {
        "text": "我的邮箱账号是多少",
        "tenant_key": "tenant",
        "app_id": "app",
        "open_id": "user",
    }

    private = asyncio.run(handle_email_command({**event, "chat_type": "p2p"}))
    group = asyncio.run(handle_email_command({**event, "chat_type": "group"}))

    assert "shxm18@example.com" in private.answer
    assert "shxm18@example.com" not in group.answer
    assert "私聊" in group.answer
    lookup.assert_awaited_with("tenant", "app", "user")


def test_one_week_ago_search_uses_one_calendar_day() -> None:
    plan = _fallback_plan("查找一周前的某封邮件", get_settings())
    target = (datetime.now(timezone(timedelta(hours=8))) - timedelta(days=7)).date().isoformat()
    assert plan.action == "search"
    assert (plan.date_start, plan.date_end) == (target, target)
    assert plan.lookback_hours >= 8 * 24


def test_search_filters_account_sender_and_subject_before_limit(monkeypatch) -> None:
    from app.email_repository import _search_messages

    calls = []

    class Connection:
        def execute(self, sql, *params):
            calls.append((sql, params))
            return type("Cursor", (), {"description": [("id",)], "fetchall": lambda self: [(42,)]})()

    @contextmanager
    def connection():
        yield Connection()

    monkeypatch.setattr("app.email_repository._open_connection", connection)
    rows = _search_messages(7, datetime(2026, 9, 14), datetime(2026, 9, 16), "张三", "项目进度")

    assert rows == [{"id": 42}]
    sql, params = calls[0]
    assert sql.index("WHERE m.email_account_id") < sql.index("ORDER BY")
    assert sql.count("CHARINDEX") == 2
    assert params[0] == 7 and params[-2:] == ("张三", "项目进度")


def test_search_then_analyze_selected_message_is_account_scoped(monkeypatch) -> None:
    account = {"id": 7, "email_address": "user@example.com", "retention_days": 14}
    lookup = AsyncMock(return_value=account)
    search = AsyncMock(return_value=[{
        "id": 42, "sent_at": datetime(2026, 9, 15, 1, 0),
        "sender_name": "张三", "subject": "项目进度",
    }])
    message = {"id": 42, "summary": "已分析", "subject": "项目进度"}
    get_message = AsyncMock(return_value=message)
    planner = AsyncMock(return_value=EmailPlan(
        action="search", lookback_hours=216, date_start="2026-09-15",
        date_end="2026-09-15", sender_contains="张三",
    ))
    monkeypatch.setattr("app.email_service.get_email_account", lookup)
    monkeypatch.setattr("app.email_service.plan_email_request", planner)
    monkeypatch.setattr("app.email_service.sync_account", AsyncMock())
    monkeypatch.setattr("app.email_service.search_messages", search)
    monkeypatch.setattr("app.email_service.get_message", get_message)
    monkeypatch.setattr("app.email_service.analyze_pending", AsyncMock())
    monkeypatch.setattr("app.email_service.create_push_logs", AsyncMock())
    monkeypatch.setattr("app.email_service.build_email_report", lambda *_, **__: "分析结果")
    monkeypatch.setattr("app.email_service.build_email_report_card", lambda *_, **__: {"body": {"elements": []}})
    event = {"tenant_key": "tenant", "app_id": "app", "open_id": "user", "chat_id": "chat"}

    found = asyncio.run(handle_email_command({**event, "text": "查找一周前张三的邮件"}))
    selected = asyncio.run(handle_email_command({**event, "text": "分析邮件 #42"}))

    assert "#42" in found.answer and "项目进度" in found.answer
    assert selected.answer == "分析结果"
    lookup.assert_awaited_with("tenant", "app", "user")
    assert get_message.await_args_list[0].args == (7, 42)
    assert planner.await_count == 1
    assert search.await_args.args[3] == "张三"
    assert "指定邮件" in build_email_report(account, [message], EmailPlan(action="query"), range_label="指定邮件")
    assert "指定邮件" in build_email_report_card(account, [message], EmailPlan(action="query"), range_label="指定邮件")["body"]["elements"][0]["content"]


def test_selected_old_email_reloads_body_without_storing_it(monkeypatch) -> None:
    account = {
        "id": 7, "email_address": "user@example.com", "retention_days": 7,
        "pop3_host": "pop.example.com", "pop3_port": 995,
        "password_ciphertext": "encrypted",
    }
    old = {"id": 42, "pop3_uidl": "wanted", "summary": None, "text_body": None, "html_body": None}
    analyzed = {**old, "summary": "旧邮件分析结果"}
    get_message = AsyncMock(side_effect=[old, analyzed])
    load = AsyncMock()
    monkeypatch.setattr("app.email_service.get_email_account", AsyncMock(return_value=account))
    monkeypatch.setattr("app.email_service.get_message", get_message)
    monkeypatch.setattr("app.email_service.decrypt_email_password", lambda _: "secret")
    monkeypatch.setattr("app.email_service.fetch_message_by_uidl", lambda **_: type(
        "Mail", (), {"to_record": lambda self: {"text_body": "原邮箱正文", "html_body": ""}}
    )())
    monkeypatch.setattr("app.email_service.analyze_pending", load)
    monkeypatch.setattr("app.email_service.create_push_logs", AsyncMock())
    monkeypatch.setattr("app.email_service.build_email_report", lambda *_, **__: "分析结果")
    monkeypatch.setattr("app.email_service.build_email_report_card", lambda *_, **__: {"body": {"elements": []}})
    event = {"tenant_key": "tenant", "app_id": "app", "open_id": "user", "chat_id": "chat", "text": "分析邮件 #42"}

    result = asyncio.run(handle_email_command(event))

    assert result.answer == "分析结果"
    assert load.await_args.args[0][0]["text_body"] == "原邮箱正文"
    assert get_message.await_args_list[0].args == (7, 42)


def test_model_recognizes_dynamic_email_account_command(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    response = type(
        "Response",
        (),
        {"tool_calls": [{"name": "plan_email_request", "args": {"action": "rebind"}}]},
    )()
    monkeypatch.setattr(
        "app.email_service.invoke_chat_with_fallback",
        AsyncMock(return_value=response),
    )

    plan = asyncio.run(plan_email_request("更换账号"))

    assert plan.action == "rebind"


def test_ambiguous_email_command_returns_specific_clarification(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    response = type(
        "Response",
        (),
        {
            "tool_calls": [
                {
                    "name": "plan_email_request",
                    "args": {
                        "action": "clarify",
                        "clarification": "请说明要分析的时间范围。",
                    },
                }
            ]
        },
    )()
    account_lookup = AsyncMock()
    monkeypatch.setattr(
        "app.email_service.invoke_chat_with_fallback", AsyncMock(return_value=response)
    )
    monkeypatch.setattr("app.email_service.get_email_account", account_lookup)

    result = asyncio.run(
        handle_email_command(
            {
                "text": "帮我处理一下邮箱里的内容",
                "tenant_key": "tenant",
                "app_id": "app",
                "open_id": "user",
            }
        )
    )

    assert result.answer == "请说明要分析的时间范围。"
    account_lookup.assert_not_awaited()


def test_unsupported_email_mutation_explains_supported_alternative() -> None:
    plan = _fallback_plan("把这封邮件转发邮件给王五", get_settings())

    assert plan.action == "clarify"
    assert "暂不支持" in str(plan.clarification)


def test_unrelated_command_returns_to_general_agent(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    response = type(
        "Response",
        (),
        {"tool_calls": [{"name": "plan_email_request", "args": {"action": "unrelated"}}]},
    )()
    account_lookup = AsyncMock()
    monkeypatch.setattr(
        "app.email_service.invoke_chat_with_fallback",
        AsyncMock(return_value=response),
    )
    monkeypatch.setattr("app.email_service.get_email_account", account_lookup)

    result = asyncio.run(
        handle_email_command(
            {
                "text": "查询生产进度",
                "tenant_key": "tenant",
                "app_id": "app",
                "open_id": "user",
            }
        )
    )

    assert result is None
    account_lookup.assert_not_awaited()


def test_explicit_numbers_override_inconsistent_model_plan() -> None:
    plan = _validated_plan(
        EmailPlan(action="query", lookback_hours=48, limit=3, only_unprocessed=False),
        get_settings(),
        "分析最近7天的前5封未处理邮件",
    )

    assert plan.lookback_hours == 168
    assert plan.limit == 5
    assert plan.only_unprocessed is True


def test_email_plan_selects_requested_messages_before_analysis() -> None:
    messages = [
        {
            "id": 1,
            "subject": "报价确认",
            "sender_name": "张三",
            "sender_address": "zhang@example.com",
            "text_body": "项目甲",
            "push_status": "pending",
        },
        {
            "id": 2,
            "subject": "普通通知",
            "sender_name": "李四",
            "sender_address": "li@example.com",
            "text_body": "项目乙",
            "push_status": "success",
        },
        {
            "id": 3,
            "subject": "报价更新",
            "sender_name": "张三",
            "sender_address": "zhang@example.com",
            "text_body": "项目甲",
            "push_status": "pending",
        },
    ]
    plan = EmailPlan(
        action="query",
        sender_contains="张三",
        keywords=["项目甲"],
        message_positions=[2],
        only_unprocessed=True,
        limit=1,
    )

    assert [item["id"] for item in _select_messages(messages, plan)] == [3]


def test_analysis_failure_is_not_reported_as_no_matching_email(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    message = {
        "id": 1,
        "to_json": '[{"address":"user@example.com"}]',
        "cc_json": "[]",
        "text_body": "",
        "html_body": "",
        "analysis_status": "pending",
        "push_status": "pending",
        "summary": None,
    }
    failed = dict(message, analysis_status="failed")
    monkeypatch.setattr("app.email_service.sync_account", AsyncMock(return_value=0))
    monkeypatch.setattr(
        "app.email_service.list_recent_messages",
        AsyncMock(side_effect=[[message], [failed]]),
    )
    monkeypatch.setattr("app.email_service._display_name", AsyncMock(return_value=""))
    monkeypatch.setattr("app.email_service.analyze_pending", AsyncMock())

    result = asyncio.run(
        sync_analyze_report(
            {"message_id": "message-1"},
            {"id": 1, "email_address": "user@example.com", "retention_days": 7},
            EmailPlan(action="query"),
        )
    )

    assert result.status == "error"
    assert "大模型分析暂时不可用" in result.answer


def test_scheduled_report_keeps_previously_pushed_email(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    message = {
        "id": 314,
        "subject": "Microsoft Outlook 测试消息",
        "sender_name": "Microsoft Outlook",
        "sender_address": "support@example.com",
        "to_json": '[{"address":"user@example.com"}]',
        "cc_json": "[]",
        "text_body": "测试邮件",
        "html_body": "",
        "analysis_status": "success",
        "summary": "账户设置测试邮件。",
        "importance": "低",
        "requires_attention": False,
        "relation_type": "To",
        "todos_json": "[]",
        "risks_json": "[]",
        "attachments_json": "[]",
    }
    monkeypatch.setattr("app.email_service.sync_account", AsyncMock(return_value=0))
    monkeypatch.setattr(
        "app.email_service.list_recent_messages",
        AsyncMock(side_effect=[[message], [message]]),
    )
    monkeypatch.setattr("app.email_service._display_name", AsyncMock(return_value=""))
    monkeypatch.setattr("app.email_service.analyze_pending", AsyncMock())
    monkeypatch.setattr("app.email_service.create_push_logs", AsyncMock())

    result = asyncio.run(
        sync_analyze_report(
            {"message_id": "scheduled:6:run", "_automation_run": True, "_email_task_id": 6},
            {"id": 1, "email_address": "user@example.com", "retention_days": 7},
            EmailPlan(action="query", lookback_hours=48),
        )
    )

    assert result.status == "success"
    assert "Microsoft Outlook 测试消息" in result.answer


def test_invalid_analysis_batch_retries_each_message(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    messages = [{"id": 1, "summary": None}, {"id": 2, "summary": None}]
    analyses = [
        EmailAnalysis(
            message_id=message_id,
            summary=f"摘要 {message_id}",
            importance="中",
            requires_attention=True,
            relation_type="To",
        )
        for message_id in (1, 2)
    ]
    analyze = AsyncMock(side_effect=[ValueError("invalid batch"), [analyses[0]], [analyses[1]]])
    save = AsyncMock()
    failed = AsyncMock()
    monkeypatch.setattr("app.email_service._analyze_batch", analyze)
    monkeypatch.setattr("app.email_service.save_analysis", save)
    monkeypatch.setattr("app.email_service.mark_analysis_failed", failed)

    asyncio.run(analyze_pending(messages, 7, "user@example.com", "User"))

    assert analyze.await_count == 3
    assert save.await_count == 2
    failed.assert_not_awaited()


def test_partial_analysis_failure_is_visible_to_user(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    base = {
        "to_json": '[{"address":"user@example.com"}]',
        "cc_json": "[]",
        "text_body": "",
        "html_body": "",
        "push_status": "pending",
    }
    pending = [
        dict(base, id=1, analysis_status="pending", summary=None),
        dict(base, id=2, analysis_status="pending", summary=None),
    ]
    refreshed = [
        dict(
            base,
            id=1,
            analysis_status="success",
            summary="需要回复客户。",
            importance="高",
            requires_attention=True,
            relation_type="To",
            todos_json="[]",
            risks_json="[]",
            attachments_json="[]",
        ),
        dict(base, id=2, analysis_status="failed", summary=None),
    ]
    monkeypatch.setattr("app.email_service.sync_account", AsyncMock(return_value=0))
    monkeypatch.setattr(
        "app.email_service.list_recent_messages",
        AsyncMock(side_effect=[pending, refreshed]),
    )
    monkeypatch.setattr("app.email_service._display_name", AsyncMock(return_value=""))
    monkeypatch.setattr("app.email_service.analyze_pending", AsyncMock())
    monkeypatch.setattr("app.email_service.create_push_logs", AsyncMock())

    result = asyncio.run(
        sync_analyze_report(
            {"message_id": "message-1"},
            {"id": 1, "email_address": "user@example.com", "retention_days": 7},
            EmailPlan(action="query"),
        )
    )

    assert result.status == "success"
    assert "1 封邮件本次分析失败" in result.answer
    assert any(
        "1 封邮件本次分析失败" in item.get("content", "")
        for item in result.card["body"]["elements"]
    )


def test_email_report_has_no_action_buttons() -> None:
    report = build_email_report(
        {"email_address": "user@example.com"},
        [
            {
                "subject": "报价确认",
                "sender_name": "供应商",
                "importance": "高",
                "relation_type": "To",
                "summary": "需要确认报价。",
                "todos_json": '["确认报价"]',
                "possible_owner": "采购部",
                "deadline": "本周五",
                "risks_json": '["逾期风险"]',
                "attachments_json": (
                    '[{"name":"quote.xlsx","analysis":"Excel，共1个工作表；报价：2行×2列"}]'
                ),
            }
        ],
        _fallback_plan("分析最近两天的重要邮件", get_settings()),
    )
    assert "确认报价" in report
    assert "quote.xlsx" in report
    assert "附件分析" in report
    assert "已处理" not in report
    assert "忽略按钮" not in report


def test_email_report_card_is_an_ai_briefing_without_table() -> None:
    card = build_email_report_card(
        {"email_address": "user@example.com"},
        [
            {
                "subject": "报价确认",
                "sender_name": "供应商",
                "importance": "高",
                "requires_attention": True,
                "relation_type": "To",
                "summary": "供应商等待报价确认，建议今天回复。",
                "todos_json": '["确认报价并回复"]',
                "possible_owner": "采购部",
                "deadline": "今天",
                "risks_json": '["逾期可能影响交付"]',
                "attachments_json": (
                    '[{"name":"quote.docx","analysis":"Word 文档内容：报价有效期为30天"}]'
                ),
            }
        ],
        EmailPlan(action="query"),
    )

    assert card["header"]["title"]["content"] == "邮件分析简报"
    assert all(item["tag"] != "table" for item in card["body"]["elements"])
    content = "\n".join(
        item.get("content", "") for item in card["body"]["elements"]
    )
    assert "分析结果" in content
    assert "附件分析" in content
    assert "报价有效期为30天" in content
    assert "确认报价并回复" in content
    assert "逾期可能影响交付" in content


def test_reply_needed_report_is_clearly_a_pop3_suggestion() -> None:
    plan = EmailPlan(action="query", reply_needed_only=True)
    card = build_email_report_card(
        {"email_address": "user@example.com"},
        [],
        plan,
    )

    assert card["header"]["title"]["content"] == "待回复邮件建议"
    content = "\n".join(
        item.get("content", "") for item in card["body"]["elements"]
    )
    assert "疑似需要你回复" in content
    assert "POP3 无法读取已发送邮件" in content


def test_reply_needed_filter_requires_an_action_and_excludes_no_reply() -> None:
    assert _likely_needs_reply(
        {
            "requires_attention": True,
            "sender_address": "supplier@example.com",
            "todos_json": '["确认报价并回复"]',
        }
    )
    assert not _likely_needs_reply(
        {
            "requires_attention": True,
            "sender_address": "no-reply@example.com",
            "todos_json": '["查看通知"]',
        }
    )
    assert not _likely_needs_reply(
        {
            "requires_attention": True,
            "sender_address": "supplier@example.com",
            "todos_json": "[]",
        }
    )


def test_email_report_card_paginates_all_details() -> None:
    messages = [
        {
            "subject": f"历史邮件 {index}",
            "sender_address": "sender@example.com",
            "importance": "中",
            "requires_attention": False,
            "relation_type": "To",
            "summary": "测试摘要",
            "todos_json": "[]",
            "risks_json": "[]",
            "attachments_json": "[]",
        }
        for index in range(1, 81)
    ]

    card = build_email_report_card(
        {"email_address": "user@example.com"},
        messages,
        EmailPlan(action="query", limit=80),
        run_ref="run-1",
    )
    content = "\n".join(item.get("content", "") for item in card["body"]["elements"])
    pager = card["body"]["elements"][-1]

    assert "共 **80** 封" in content
    assert "第 **1/4** 页" in content
    assert "历史邮件 1" in content
    assert "历史邮件 21" not in content
    assert "87600" not in content
    assert pager["tag"] == "form"
    assert pager["elements"][0]["form_action_type"] == "submit"
    assert pager["elements"][0]["behaviors"][0]["value"] == {
        "action": "email_page",
        "run_ref": "run-1",
        "page": 2,
    }

    second = build_email_report_card(
        {"email_address": "user@example.com"},
        messages,
        EmailPlan(action="query", limit=80),
        page=2,
        run_ref="run-1",
    )
    second_content = "\n".join(
        item.get("content", "") for item in second["body"]["elements"]
    )
    assert "第 **2/4** 页" in second_content
    assert "历史邮件 21" in second_content
    assert len(second["body"]["elements"][-1]["elements"]) == 2


def test_email_report_card_groups_direct_and_cc_messages() -> None:
    messages = [
        {
            "subject": f"邮件 {index}",
            "sender_address": "sender@example.com",
            "importance": "中",
            "requires_attention": False,
            "relation_type": "Cc" if index <= 10 else "To",
            "summary": "测试摘要",
            "todos_json": "[]",
            "risks_json": "[]",
            "attachments_json": "[]",
        }
        for index in range(1, 22)
    ]
    plan = EmailPlan(
        action="query", scope="to_or_cc", group_by_recipient=True, limit=100
    )

    card = build_email_report_card(
        {"email_address": "user@example.com"}, messages, plan, run_ref="run-1"
    )
    content = "\n".join(item.get("content", "") for item in card["body"]["elements"])
    button_value = card["body"]["elements"][-1]["elements"][0]["behaviors"][0][
        "value"
    ]

    assert content.index("直接发给我的") < content.index("抄送给我的")
    assert button_value["group_by_recipient"] is True


def test_bind_card_uses_password_input_and_show_toggle() -> None:
    card = build_email_login_card("绑定公司邮箱", "one-time-token")
    form = card["body"]["elements"][1]
    password = form["elements"][1]
    submit = form["elements"][2]

    assert card["config"]["enable_forward"] is False
    assert password["input_type"] == "password"
    assert password["show_icon"] is True
    assert submit["form_action_type"] == "submit"
    assert submit["behaviors"][0]["value"]["token"] == "one-time-token"


def test_extract_email_bind_card_action_keeps_tenant_and_user_scope() -> None:
    app = TenantApp("tenant-1", "app-1", "secret", None, None, "bot-1", None)
    payload = {
        "event": {
            "operator": {"tenant_key": "tenant-1", "open_id": "user-1"},
            "context": {"open_chat_id": "chat-1", "open_message_id": "message-1"},
            "action": {
                "tag": "button",
                "value": {"action": "email_bind", "token": "token-1"},
                "form_value": {
                    "email_account": "user@example.com",
                    "email_password": "secret",
                },
            },
        }
    }

    action = extract_email_bind_card_action(payload, app)

    assert action == {
        "tenant_key": "tenant-1",
        "app_id": "app-1",
        "bot_code": "bot-1",
        "open_id": "user-1",
        "chat_id": "chat-1",
        "message_id": "message-1",
        "token": "token-1",
        "email": "user@example.com",
        "password": "secret",
    }


def test_extract_email_bind_card_action_rejects_other_tenant() -> None:
    app = TenantApp("tenant-1", "app-1", "secret", None, None, "bot-1", None)
    payload = {
        "event": {
            "operator": {"tenant_key": "tenant-2", "open_id": "user-1"},
            "context": {"open_chat_id": "chat-1", "open_message_id": "message-1"},
            "action": {
                "tag": "button",
                "value": {"action": "email_bind", "token": "token-1"},
                "form_value": {
                    "email_account": "user@example.com",
                    "email_password": "secret",
                },
            },
        }
    }

    assert extract_email_bind_card_action(payload, app) is None


def test_extract_email_page_action_keeps_user_scope() -> None:
    app = TenantApp("tenant-1", "app-1", "secret", None, None, "bot-1", None)
    payload = {
        "event": {
            "operator": {"tenant_key": "tenant-1", "open_id": "user-1"},
            "context": {"open_chat_id": "chat-1", "open_message_id": "message-1"},
            "action": {
                "tag": "button",
                "value": {
                    "action": "email_page",
                    "run_ref": "run-1",
                    "page": 2,
                    "group_by_recipient": True,
                },
            },
        }
    }

    assert extract_email_page_card_action(payload, app) == {
        "tenant_key": "tenant-1",
        "app_id": "app-1",
        "open_id": "user-1",
        "chat_id": "chat-1",
        "message_id": "message-1",
        "run_ref": "run-1",
        "page": 2,
        "group_by_recipient": True,
    }


def test_email_page_callback_updates_original_card(monkeypatch: pytest.MonkeyPatch) -> None:
    app = TenantApp("tenant-1", "app-1", "secret", None, None, "bot-1", None)
    card = {"body": {"elements": []}}
    render = AsyncMock(return_value=card)
    update = AsyncMock(return_value=True)
    monkeypatch.setattr("app.feishu_ws.render_email_report_page", render)
    monkeypatch.setattr("app.feishu_ws.update_card", update)

    asyncio.run(
        process_email_page_card(
            app,
            {
                "tenant_key": "tenant-1",
                "app_id": "app-1",
                "open_id": "user-1",
                "chat_id": "chat-1",
                "message_id": "message-1",
                "run_ref": "run-1",
                "page": 2,
                "group_by_recipient": True,
            },
        )
    )

    render.assert_awaited_once_with(
        "tenant-1",
        "app-1",
        "user-1",
        "run-1",
        2,
        group_by_recipient=True,
    )
    update.assert_awaited_once_with(app, "message-1", card)


def test_card_callback_frames_reach_the_sdk_dispatcher() -> None:
    seen: list[bytes] = []

    class Handler:
        def _do_without_validation(self, payload: bytes) -> None:
            seen.append(payload)

    client = object.__new__(CardCallbackWsClient)
    client._event_handler = Handler()
    client._write_message = AsyncMock()
    frame = Frame(SeqID=1, LogID=1, service=1, method=1, payload=b'{"schema":"2.0"}')
    for key, value in ((HEADER_TYPE, "card"), (HEADER_SUM, "1"), (HEADER_SEQ, "0")):
        header = frame.headers.add()
        header.key = key
        header.value = value

    asyncio.run(client._handle_data_frame(frame))

    assert seen == [b'{"schema":"2.0"}']
    client._write_message.assert_awaited_once()


def test_successful_card_binding_consumes_token_and_updates_card(monkeypatch: pytest.MonkeyPatch) -> None:
    app = TenantApp("tenant-1", "app-1", "secret", None, None, "bot-1", None)
    action = {
        "tenant_key": "tenant-1",
        "app_id": "app-1",
        "bot_code": "bot-1",
        "open_id": "user-1",
        "chat_id": "chat-1",
        "message_id": "message-1",
        "token": "token-1",
        "email": "user@example.com",
        "password": "secret",
    }
    scope = {key: action[key] for key in ("tenant_key", "app_id", "bot_code", "open_id", "chat_id")}
    account = {"id": 1}
    bind = AsyncMock(return_value=account)
    consume = AsyncMock()
    sync = AsyncMock(return_value=3)
    update = AsyncMock(return_value=True)
    send = AsyncMock()
    monkeypatch.setattr("app.feishu_ws.get_bind_scope", AsyncMock(return_value=scope))
    monkeypatch.setattr("app.feishu_ws.bind_email_account", bind)
    monkeypatch.setattr("app.feishu_ws.consume_bind_token", consume)
    monkeypatch.setattr("app.feishu_ws.initial_sync", sync)
    monkeypatch.setattr("app.feishu_ws.update_card", update)
    monkeypatch.setattr("app.feishu_ws.send_card", send)

    asyncio.run(process_email_bind_card(app, action))

    bind.assert_awaited_once()
    assert bind.await_args.args[1:] == ("user@example.com", "secret")
    consume.assert_awaited_once_with("token-1")
    sync.assert_awaited_once_with(account)
    assert "首次同步新增 3 封邮件" in update.await_args.args[2]["body"]["elements"][0]["content"]
    assert update.await_args.args[2]["header"]["title"]["content"] == "登录成功"
    send.assert_awaited_once()
    assert send.await_args.args[2]["header"]["title"]["content"] == "功能使用介绍"
    assert "邮箱已绑定" in str(send.await_args.args[2])


def test_wrong_card_password_keeps_token_for_retry(monkeypatch: pytest.MonkeyPatch) -> None:
    app = TenantApp("tenant-1", "app-1", "secret", None, None, "bot-1", None)
    action = {
        "tenant_key": "tenant-1",
        "app_id": "app-1",
        "bot_code": "bot-1",
        "open_id": "user-1",
        "chat_id": "chat-1",
        "message_id": "message-1",
        "token": "token-1",
        "email": "user@example.com",
        "password": "wrong",
    }
    scope = {key: action[key] for key in ("tenant_key", "app_id", "bot_code", "open_id", "chat_id")}
    consume = AsyncMock()
    send = AsyncMock()
    monkeypatch.setattr("app.feishu_ws.get_bind_scope", AsyncMock(return_value=scope))
    monkeypatch.setattr(
        "app.feishu_ws.bind_email_account",
        AsyncMock(side_effect=EmailAuthenticationError("authentication failed")),
    )
    monkeypatch.setattr("app.feishu_ws.consume_bind_token", consume)
    monkeypatch.setattr("app.feishu_ws.send_card", send)

    asyncio.run(process_email_bind_card(app, action))

    consume.assert_not_awaited()
    assert "邮箱账号或密码错误" in send.await_args.args[2]["body"]["elements"][0]["content"]
    send.assert_awaited_once()


def test_password_fields_are_redacted_from_logs() -> None:
    assert "secret" not in redact_sensitive_text('password=secret')
    assert "secret" not in redact_sensitive_text('{"email_password":"secret"}')
    assert "part-two" not in redact_sensitive_text("UID=user;PWD={secret;part-two};Encrypt=yes")
