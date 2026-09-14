from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from app.bootstrap import bootstrap
from app.excel_design import (
    ExcelAggregation,
    ExcelDeduplication,
    ExcelExportPlan,
    ExcelOutputColumn,
    ExcelSort,
    apply_excel_export_plan,
    plan_excel_export,
)
from app.excel_export import export_latest_result_to_excel
from app.export_context import save_export_context
from app.identity import Identity
from app.policy import PolicyResult
from app.routing.export_intent import (
    is_explicit_excel_export_request,
    parse_export_intent_response,
    should_export_excel,
)
from app.settings import get_settings
from app.tool_result_cache import save_tool_result


class ExcelExportIntentTests(unittest.IsolatedAsyncioTestCase):
    def test_common_excel_typo_is_recognized(self) -> None:
        self.assertTrue(is_explicit_excel_export_request("整理成 excle 表给我"))

    def test_natural_export_requests_are_recognized(self) -> None:
        requests = (
            "导出 Excel",
            "帮我整理成excel表给我",
            "把刚才的结果生成 Excel 文件",
            "可以做成Excel发我吗",
        )
        for text in requests:
            with self.subTest(text=text):
                self.assertTrue(is_explicit_excel_export_request(text))

    def test_excel_analysis_is_not_mistaken_for_export(self) -> None:
        self.assertFalse(is_explicit_excel_export_request("帮我分析这个 Excel 文件"))

    async def test_ambiguous_spreadsheet_request_uses_model_fallback(self) -> None:
        response = type(
            "Response",
            (),
            {
                "content": json.dumps(
                    {
                        "intent": "export_excel",
                        "confidence": 0.96,
                        "reason": "用户要求把已有结果交付为电子表格",
                    },
                    ensure_ascii=False,
                )
            },
        )()
        with patch(
            "app.routing.export_intent.invoke_chat_with_fallback",
            AsyncMock(return_value=response),
        ) as classifier:
            result = await should_export_excel("把上面的数据弄个电子表格")

        self.assertTrue(result)
        classifier.assert_awaited_once()

    async def test_low_confidence_model_result_does_not_export(self) -> None:
        response = type(
            "Response",
            (),
            {"content": '{"intent":"export_excel","confidence":0.55}'},
        )()
        with patch(
            "app.routing.export_intent.invoke_chat_with_fallback",
            AsyncMock(return_value=response),
        ):
            self.assertFalse(await should_export_excel("帮我生成一份业务报表"))

    def test_invalid_classifier_output_is_safe(self) -> None:
        decision = parse_export_intent_response("不是 JSON")
        self.assertEqual(decision.intent, "normal")
        self.assertEqual(decision.confidence, 0.0)


class SessionScopedExportTests(unittest.IsolatedAsyncioTestCase):
    async def test_model_designs_deduplication_using_real_dataset_column(self) -> None:
        response = type(
            "Response",
            (),
            {
                "content": "",
                "tool_calls": [
                    {
                        "name": "design_excel_export",
                        "args": {
                            "sheet_name": "去重结果",
                            "title": "",
                            "filename_suffix": "样品号去重",
                            "filters": [],
                            "deduplicate": {
                                "keys": ["样品单号"],
                                "keep": "merge_unique",
                                "include_count": True,
                            },
                            "group_by": [],
                            "aggregations": [],
                            "sort_by": [],
                            "columns": [],
                        },
                    }
                ],
            },
        )()
        rows = [{"样品单号": "S001", "节点": "裁剪"}]
        with patch(
            "app.excel_design.invoke_chat_with_fallback",
            AsyncMock(return_value=response),
        ) as model:
            plan = await plan_excel_export("把样品的号去重后生成 Excel", rows)

        self.assertEqual(plan.deduplicate.keys, ["样品单号"])
        request_payload = model.await_args.kwargs["messages"][1].content
        self.assertIn("样品单号", request_payload)

    def test_generic_deduplication_merges_distinct_values(self) -> None:
        rows = [
            {"样品单号": "S001", "客户简称": "NIKE", "节点": "打样 / 裁剪"},
            {"样品单号": "S001", "客户简称": "NIKE", "节点": "打样 / 车缝"},
            {"样品单号": "S002", "客户简称": "VUO", "节点": "确认 / 包装"},
        ]
        plan = ExcelExportPlan(
            deduplicate=ExcelDeduplication(
                keys=["样品单号"], keep="merge_unique", include_count=True
            )
        )

        result = apply_excel_export_plan(rows, plan)

        self.assertEqual(len(result), 2)
        self.assertEqual(result[0]["样品单号"], "S001")
        self.assertEqual(result[0]["客户简称"], "NIKE")
        self.assertEqual(result[0]["节点"], "打样 / 裁剪；打样 / 车缝")
        self.assertEqual(result[0]["合并记录数"], 2)

    def test_generic_group_sort_and_column_design(self) -> None:
        rows = [
            {"客户": "B", "金额": 5, "单号": "B1"},
            {"客户": "A", "金额": 10, "单号": "A1"},
            {"客户": "A", "金额": 20, "单号": "A2"},
        ]
        plan = ExcelExportPlan(
            group_by=["客户"],
            aggregations=[
                ExcelAggregation(column="金额", function="sum", header="总金额"),
                ExcelAggregation(column="单号", function="count_distinct", header="订单数"),
            ],
            sort_by=[ExcelSort(column="总金额", direction="desc")],
            columns=[
                ExcelOutputColumn(source="客户", header="客户名称"),
                ExcelOutputColumn(source="总金额"),
                ExcelOutputColumn(source="订单数"),
            ],
        )

        result = apply_excel_export_plan(rows, plan)

        self.assertEqual(
            result,
            [
                {"客户名称": "A", "总金额": 30, "订单数": 2},
                {"客户名称": "B", "总金额": 5, "订单数": 1},
            ],
        )

    async def asyncSetUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(dir=Path.cwd())
        self.root = Path(self.temp_dir.name)
        self.previous = {
            key: os.environ.get(key)
            for key in ("APP_DATABASE_PATH", "FEISHU_APPS_CONFIG_PATH", "APP_ENV")
        }
        os.environ["APP_DATABASE_PATH"] = str(self.root / "agent.db")
        os.environ["FEISHU_APPS_CONFIG_PATH"] = str(self.root / "missing-apps.json")
        os.environ["APP_ENV"] = "production"
        get_settings.cache_clear()
        await bootstrap()

    async def asyncTearDown(self) -> None:
        get_settings.cache_clear()
        for key, value in self.previous.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        self.temp_dir.cleanup()

    async def test_current_file_analysis_replaces_older_tool_export(self) -> None:
        scope = {
            "tenant_key": "tenant-a",
            "app_id": "app-a",
            "open_id": "user-a",
            "chat_id": "chat-a",
            "session_id": "session-a",
        }
        old_result_id = await save_tool_result(
            request_id="request-old",
            tenant_key=scope["tenant_key"],
            app_id=scope["app_id"],
            open_id=scope["open_id"],
            bot_code="bot",
            message_id="message-old",
            chat_id=scope["chat_id"],
            tool_name="sampleschedule_risk",
            tool_args={},
            tool_result=json.dumps({"rows": [{"旧结果": "不应导出"}]}, ensure_ascii=False),
        )
        await save_export_context(
            **scope,
            source_type="tool_result",
            source_ref=str(old_result_id),
            source_name="sampleschedule_risk",
        )
        analysis = """## 客户贡献
| 客户 | 销售额 |
|---|---:|
| AG | 986744 |
| BG | 200000 |
"""
        await save_export_context(
            **scope,
            source_type="analysis_result",
            source_ref=json.dumps(
                {"answer": analysis, "source_path": str(self.root / "sales.xlsx")},
                ensure_ascii=False,
            ),
            source_name="2026年销售数据.xlsx",
        )

        with patch("app.excel_export.EXPORT_DIR", self.root / "exports"):
            path = await export_latest_result_to_excel(**scope)

        self.assertIsNotNone(path)
        workbook = load_workbook(path, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        self.assertEqual(sheet["A1"].value, "客户")
        self.assertEqual(sheet["A2"].value, "AG")
        self.assertEqual(sheet["B2"].value, 986744)
        self.assertNotEqual(sheet["A1"].value, "旧结果")

    async def test_new_session_cannot_export_previous_session_context(self) -> None:
        await save_export_context(
            tenant_key="tenant-a",
            app_id="app-a",
            open_id="user-a",
            chat_id="chat-a",
            session_id="old-session",
            source_type="uploaded_file",
            source_ref=str(self.root / "old.xlsx"),
            source_name="old.xlsx",
        )

        path = await export_latest_result_to_excel(
            tenant_key="tenant-a",
            app_id="app-a",
            open_id="user-a",
            chat_id="chat-a",
            session_id="new-session",
        )

        self.assertIsNone(path)

    async def test_reply_exports_the_referenced_result_instead_of_latest_result(self) -> None:
        scope = {
            "tenant_key": "tenant-a",
            "app_id": "app-a",
            "open_id": "user-a",
            "chat_id": "chat-a",
            "session_id": "session-a",
        }
        old_result_id = await save_tool_result(
            request_id="request-old",
            tenant_key=scope["tenant_key"],
            app_id=scope["app_id"],
            open_id=scope["open_id"],
            bot_code="bot",
            message_id="user-query-old",
            chat_id=scope["chat_id"],
            tool_name="sample_advice_notice",
            tool_args={},
            tool_result=json.dumps({"rows": [{"样品单号": "OLD"}]}, ensure_ascii=False),
        )
        await save_export_context(
            **scope,
            source_type="tool_result",
            source_ref=str(old_result_id),
            source_name="sample_advice_notice",
            request_message_id="user-query-old",
            reply_message_id="bot-reply-old",
        )
        new_result_id = await save_tool_result(
            request_id="request-new",
            tenant_key=scope["tenant_key"],
            app_id=scope["app_id"],
            open_id=scope["open_id"],
            bot_code="bot",
            message_id="user-query-new",
            chat_id=scope["chat_id"],
            tool_name="sample_advice_notice",
            tool_args={},
            tool_result=json.dumps({"rows": [{"样品单号": "NEW"}]}, ensure_ascii=False),
        )
        await save_export_context(
            **scope,
            source_type="tool_result",
            source_ref=str(new_result_id),
            source_name="sample_advice_notice",
            request_message_id="user-query-new",
            reply_message_id="bot-reply-new",
        )

        with patch("app.excel_export.EXPORT_DIR", self.root / "exports"):
            path = await export_latest_result_to_excel(
                **scope,
                referenced_message_ids=["bot-reply-old"],
            )

        self.assertIsNotNone(path)
        workbook = load_workbook(path, data_only=True)
        self.assertEqual(workbook.active["A2"].value, "OLD")

    async def test_unknown_reply_does_not_fall_back_to_latest_result(self) -> None:
        scope = {
            "tenant_key": "tenant-a",
            "app_id": "app-a",
            "open_id": "user-a",
            "chat_id": "chat-a",
            "session_id": "session-a",
        }
        result_id = await save_tool_result(
            request_id="request-new",
            tenant_key=scope["tenant_key"],
            app_id=scope["app_id"],
            open_id=scope["open_id"],
            bot_code="bot",
            message_id="user-query-new",
            chat_id=scope["chat_id"],
            tool_name="sample_advice_notice",
            tool_args={},
            tool_result=json.dumps({"rows": [{"样品单号": "NEW"}]}, ensure_ascii=False),
        )
        await save_export_context(
            **scope,
            source_type="tool_result",
            source_ref=str(result_id),
            source_name="sample_advice_notice",
            request_message_id="user-query-new",
            reply_message_id="bot-reply-new",
        )

        path = await export_latest_result_to_excel(
            **scope,
            referenced_message_ids=["unrelated-message"],
        )

        self.assertIsNone(path)

    async def test_chat_scope_does_not_leak_when_chat_id_is_missing(self) -> None:
        await save_export_context(
            tenant_key="tenant-a",
            app_id="app-a",
            open_id="user-a",
            chat_id="group-chat",
            session_id="session-a",
            source_type="uploaded_file",
            source_ref=str(self.root / "group.xlsx"),
            source_name="group.xlsx",
        )

        path = await export_latest_result_to_excel(
            tenant_key="tenant-a",
            app_id="app-a",
            open_id="user-a",
            chat_id=None,
            session_id="session-a",
        )

        self.assertIsNone(path)

    async def test_remote_paginated_result_exports_all_pages(self) -> None:
        scope = {
            "tenant_key": "tenant-a",
            "app_id": "app-a",
            "open_id": "user-a",
            "chat_id": "chat-a",
            "session_id": "session-a",
        }
        first_page = {
            "rows": [{"编号": 1}, {"编号": 2}],
            "page": 1,
            "pageSize": 2,
            "totalCount": 5,
            "totalPages": 3,
        }
        result_id = await save_tool_result(
            request_id="request-remote",
            tenant_key=scope["tenant_key"],
            app_id=scope["app_id"],
            open_id=scope["open_id"],
            bot_code="bot",
            message_id="message-remote",
            chat_id=scope["chat_id"],
            tool_name="sampleshedule_risk",
            tool_args={"page": 1, "pageSize": 2},
            tool_result=json.dumps(first_page, ensure_ascii=False),
        )
        await save_export_context(
            **scope,
            source_type="tool_result",
            source_ref=str(result_id),
            source_name="sampleshedule_risk",
        )
        identity = Identity(tenant_key="tenant-a", app_id="app-a", open_id="user-a")

        async def remote_page(_client, _tool_name, arguments):
            page = arguments["page"]
            start = (page - 1) * 2 + 1
            values = list(range(start, min(start + 2, 6)))
            return json.dumps(
                {
                    "rows": [{"编号": value} for value in values],
                    "page": page,
                    "pageSize": 2,
                    "totalCount": 5,
                    "totalPages": 3,
                },
                ensure_ascii=False,
            )

        with (
            patch(
                "app.excel_export.resolve_identity", AsyncMock(return_value=identity)
            ),
            patch(
                "app.excel_export.check_agent_access",
                return_value=PolicyResult(True, "ok"),
            ),
            patch(
                "app.excel_export.check_tool_access",
                return_value=PolicyResult(True, "ok"),
            ),
            patch(
                "app.excel_export.McpClient.call_tool",
                new=remote_page,
            ),
            patch("app.excel_export.EXPORT_DIR", self.root / "exports"),
        ):
            path = await export_latest_result_to_excel(**scope)

        self.assertIsNotNone(path)
        workbook = load_workbook(path, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        self.assertEqual(sheet.max_row, 6)
        values = [sheet.cell(row=row, column=1).value for row in range(2, 7)]
        self.assertEqual(values, [1, 2, 3, 4, 5])


if __name__ == "__main__":
    unittest.main()
