from __future__ import annotations

import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from app.business_pagination import extract_pagination
from app.export_context import get_latest_export_context
from app.identity import resolve_identity
from app.mcp_client import McpClient
from app.policy import check_agent_access, check_tool_access, inject_identity_args
from app.tool_result_cache import get_tool_result_by_id


EXPORT_DIR = Path("data/exports")
logger = logging.getLogger("excel_export")


async def export_latest_result_to_excel(
    *,
    tenant_key: str,
    app_id: str,
    open_id: str,
    chat_id: str | None,
    session_id: str,
) -> Path | None:
    context = await get_latest_export_context(
        tenant_key=tenant_key,
        app_id=app_id,
        open_id=open_id,
        chat_id=chat_id,
        session_id=session_id,
    )
    if not context:
        return None

    if context["source_type"] == "uploaded_file":
        path = Path(context["source_ref"])
        return path if path.exists() and path.suffix.lower() == ".xlsx" else None

    if context["source_type"] == "analysis_result":
        return export_analysis_result(context)

    if context["source_type"] != "tool_result":
        return None

    try:
        result_id = int(context["source_ref"])
    except (TypeError, ValueError):
        return None
    cached = await get_tool_result_by_id(
        result_id=result_id,
        tenant_key=tenant_key,
        app_id=app_id,
        open_id=open_id,
        chat_id=chat_id,
    )
    if not cached:
        return None

    rows = await load_complete_tool_rows(cached)
    if not rows:
        return None

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{cached.get('tool_name') or 'query_result'}_{timestamp}.xlsx"
    path = EXPORT_DIR / safe_filename(filename)
    write_rows_to_xlsx(rows, path)
    return path


async def load_complete_tool_rows(cached: dict[str, Any]) -> list[dict[str, Any]]:
    tool_result = cached.get("tool_result") or ""
    initial_rows = extract_rows(tool_result)
    pagination = extract_pagination(tool_result)
    if not initial_rows or not pagination or pagination["totalPages"] <= 1:
        return initial_rows

    identity = await resolve_identity(
        tenant_key=str(cached["tenant_key"]),
        app_id=str(cached["app_id"]),
        open_id=str(cached["open_id"]),
        union_id=None,
        user_id=None,
    )
    agent_policy = check_agent_access(identity)
    if not agent_policy.allowed:
        logger.warning("Remote export denied: %s", agent_policy.reason)
        return []

    try:
        base_args = json.loads(cached.get("tool_args") or "{}")
    except (TypeError, json.JSONDecodeError):
        base_args = {}
    if not isinstance(base_args, dict):
        base_args = {}
    protected_args = inject_identity_args(identity, base_args)
    tool_name = str(cached["tool_name"])
    tool_policy = check_tool_access(identity, tool_name, protected_args)
    if not tool_policy.allowed:
        logger.warning("Remote export tool denied: %s", tool_policy.reason)
        return []

    rows: list[dict[str, Any]] = []
    client = McpClient()
    try:
        for page in range(1, pagination["totalPages"] + 1):
            if page == pagination["page"]:
                page_rows = initial_rows
            else:
                page_args = dict(protected_args)
                page_args["page"] = page
                page_args["pageSize"] = pagination["pageSize"]
                page_result = await client.call_tool(tool_name, page_args)
                page_meta = extract_pagination(page_result)
                page_rows = extract_rows(page_result)
                if not page_meta or page_meta["page"] != page:
                    raise ValueError(f"MCP export returned invalid page {page}")
            rows.extend(page_rows)
    except Exception:  # noqa: BLE001
        logger.exception("Remote MCP export failed: tool=%s", tool_name)
        return []
    return rows[: pagination["totalCount"]]


def export_analysis_result(context: dict[str, Any]) -> Path | None:
    try:
        payload = json.loads(context["source_ref"])
    except (KeyError, TypeError, json.JSONDecodeError):
        return None

    answer = str(payload.get("answer") or "")
    tables = extract_markdown_tables(answer)
    if not tables:
        source_path = Path(str(payload.get("source_path") or ""))
        if source_path.exists() and source_path.suffix.lower() == ".xlsx":
            return source_path
        return None

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    source_name = Path(str(context.get("source_name") or "文件分析")).stem
    path = EXPORT_DIR / safe_filename(f"{source_name}_分析结果_{timestamp}.xlsx")
    write_markdown_tables_to_xlsx(tables, path)
    return path


def extract_markdown_tables(text: str) -> list[tuple[str, list[list[str]]]]:
    lines = text.splitlines()
    tables: list[tuple[str, list[list[str]]]] = []
    pending_title = "分析结果"
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("#"):
            pending_title = line.lstrip("#").strip() or pending_title
            index += 1
            continue
        if "|" not in line or index + 1 >= len(lines):
            index += 1
            continue
        separator = lines[index + 1].strip()
        if not is_markdown_separator(separator):
            index += 1
            continue
        rows = [split_markdown_row(line)]
        index += 2
        while index < len(lines) and "|" in lines[index]:
            rows.append(split_markdown_row(lines[index]))
            index += 1
        tables.append((pending_title, rows))
    return tables


def is_markdown_separator(line: str) -> bool:
    cells = split_markdown_row(line)
    return bool(cells) and all(cell.replace(":", "").strip("-").strip() == "" and "-" in cell for cell in cells)


def split_markdown_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def write_markdown_tables_to_xlsx(
    tables: list[tuple[str, list[list[str]]]], path: Path
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()
    for number, (title, rows) in enumerate(tables, start=1):
        sheet_name = unique_sheet_name(title, number, used_names)
        ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append([normalize_markdown_cell(value) for value in row])
        for column_cells in ws.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = min(
                max(length + 2, 10), 40
            )
    wb.save(path)


def unique_sheet_name(title: str, number: int, used_names: set[str]) -> str:
    base = safe_filename(title).replace("[", "_").replace("]", "_")[:31]
    base = base or f"分析结果{number}"
    name = base
    suffix = 2
    while name in used_names:
        marker = f"_{suffix}"
        name = f"{base[: 31 - len(marker)]}{marker}"
        suffix += 1
    used_names.add(name)
    return name


def normalize_markdown_cell(value: str) -> Any:
    text = value.replace("**", "").strip()
    numeric = text.replace(",", "")
    try:
        return float(numeric) if "." in numeric else int(numeric)
    except ValueError:
        return text


def extract_rows(tool_result: str) -> list[dict[str, Any]]:
    try:
        payload = json.loads(tool_result)
    except (TypeError, json.JSONDecodeError):
        return []
    rows = payload.get("rows") if isinstance(payload, dict) else None
    if not isinstance(rows, list):
        return []
    return [row for row in rows if isinstance(row, dict)]


def write_rows_to_xlsx(rows: list[dict[str, Any]], path: Path) -> None:
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(str(key))

    wb = Workbook()
    ws = wb.active
    ws.title = "查询结果"
    ws.append(headers)
    for row in rows:
        ws.append([normalize_cell(row.get(header)) for header in headers])

    for column_cells in ws.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 40)
    wb.save(path)


def normalize_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False)


def safe_filename(name: str) -> str:
    return "".join("_" if ch in '\\/:*?"<>|' else ch for ch in name)
