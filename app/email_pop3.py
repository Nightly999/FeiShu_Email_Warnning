from __future__ import annotations

import csv
import io
import poplib
import ssl
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, timedelta, timezone
from email import policy
from email.message import Message
from email.parser import BytesParser
from email.utils import getaddresses, parsedate_to_datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook


# Some corporate mail systems return minified HTML as one long POP3 line.
# Python's 2 KiB default rejects otherwise valid messages before parsing them.
poplib._MAXLINE = 10 * 1024 * 1024
# ponytail: bound inline attachment work; move large/many files to a background worker if needed.
MAX_ANALYZED_ATTACHMENTS = 3
MAX_ATTACHMENT_BYTES = 5 * 1024 * 1024
MAX_ATTACHMENT_ANALYSIS_CHARS = 1500
MAX_UNPACKED_BYTES = 20 * 1024 * 1024


class EmailAuthenticationError(RuntimeError):
    pass


class EmailConnectionError(RuntimeError):
    pass


@dataclass(frozen=True)
class ParsedEmail:
    pop3_uidl: str
    message_id: str
    references_header: str
    in_reply_to: str
    subject: str
    sender_name: str
    sender_address: str
    to: list[dict[str, str]]
    cc: list[dict[str, str]]
    sent_at: datetime | None
    received_at: datetime
    text_body: str
    html_body: str
    attachments: list[dict[str, Any]]

    def to_record(self) -> dict[str, Any]:
        return asdict(self)


def verify_pop3_login(host: str, port: int, username: str, password: str, timeout: int) -> None:
    client = _login(host, port, username, password, timeout)
    try:
        client.noop()
    finally:
        _quit(client)


def count_pop3_messages(*, host: str, port: int, username: str, password: str, timeout: int) -> int:
    client = _login(host, port, username, password, timeout)
    try:
        return client.stat()[0]
    finally:
        _quit(client)


def fetch_message_by_uidl(
    *, host: str, port: int, username: str, password: str,
    timeout: int, uidl: str, max_body_chars: int,
) -> ParsedEmail | None:
    client = _login(host, port, username, password, timeout)
    try:
        try:
            response, lines, _ = client.uidl()
        except poplib.error_proto as exc:
            raise EmailConnectionError("邮箱服务器不支持 UIDL") from exc
        if not response.startswith(b"+OK"):
            raise EmailConnectionError("邮箱服务器不支持 UIDL")
        for line in lines:
            number, found_uidl = _uidl_entry(line)
            if found_uidl == uidl:
                response, message_lines, _ = client.retr(number)
                if response.startswith(b"+OK"):
                    return parse_message(uidl, b"\r\n".join(message_lines), max_body_chars)
                break
        return None
    finally:
        _quit(client)


def fetch_recent_messages(
    *,
    host: str,
    port: int,
    username: str,
    password: str,
    timeout: int,
    lookback_hours: int,
    max_messages: int,
    max_body_chars: int,
    known_uidls: set[str] | None = None,
    stop_at_known: bool = False,
) -> list[ParsedEmail]:
    client = _login(host, port, username, password, timeout)
    known_uidls = known_uidls or set()
    cutoff = (datetime.now(timezone.utc) - timedelta(hours=max(1, lookback_hours))).replace(
        tzinfo=None
    )
    result: list[ParsedEmail] = []
    try:
        try:
            response, lines, _ = client.uidl()
        except poplib.error_proto as exc:
            raise EmailConnectionError("邮箱服务器不支持 UIDL") from exc
        if not response.startswith(b"+OK"):
            raise EmailConnectionError("POP3 server does not support UIDL")
        entries = [_uidl_entry(line) for line in lines]
        for number, uidl in reversed(entries):
            if uidl in known_uidls:
                if stop_at_known:
                    break
                continue
            response, message_lines, _ = client.retr(number)
            if not response.startswith(b"+OK"):
                continue
            parsed = parse_message(uidl, b"\r\n".join(message_lines), max_body_chars)
            if parsed.sent_at and parsed.sent_at < cutoff:
                break
            result.append(parsed)
            if len(result) >= max(1, max_messages):
                break
        return result
    finally:
        _quit(client)


def parse_message(uidl: str, raw: bytes, max_body_chars: int) -> ParsedEmail:
    message = BytesParser(policy=policy.default).parsebytes(raw)
    sender = getaddresses([str(message.get("From") or "")])
    sent_at = _message_date(message)
    text_parts: list[str] = []
    html_parts: list[str] = []
    attachments: list[dict[str, Any]] = []

    for part in message.walk():
        if part.is_multipart():
            continue
        disposition = part.get_content_disposition()
        filename = part.get_filename()
        payload = part.get_payload(decode=True) or b""
        if disposition == "attachment" or filename:
            name = str(filename or "未命名附件")
            attachment = {
                "name": name,
                "mime_type": part.get_content_type(),
                "size": len(payload),
            }
            if len(attachments) < MAX_ANALYZED_ATTACHMENTS:
                analysis = _analyze_attachment(name, payload)
                if analysis:
                    attachment["analysis"] = analysis
            attachments.append(attachment)
            continue
        content = _part_text(part, payload)
        if part.get_content_type() == "text/plain":
            text_parts.append(content)
        elif part.get_content_type() == "text/html":
            html_parts.append(content)

    limit = max(1, max_body_chars)
    return ParsedEmail(
        pop3_uidl=uidl,
        message_id=str(message.get("Message-ID") or "")[:1000],
        references_header=str(message.get("References") or ""),
        in_reply_to=str(message.get("In-Reply-To") or "")[:1000],
        subject=str(message.get("Subject") or "")[:1000],
        sender_name=(sender[0][0] if sender else "")[:500],
        sender_address=(sender[0][1] if sender else "")[:320],
        to=_addresses(message, "To"),
        cc=_addresses(message, "Cc"),
        sent_at=sent_at,
        received_at=datetime.now(timezone.utc).replace(tzinfo=None),
        text_body="\n".join(text_parts)[:limit],
        html_body="\n".join(html_parts)[:limit],
        attachments=attachments,
    )


def _analyze_attachment(filename: str, payload: bytes) -> str | None:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".docx", ".csv", ".txt", ".md", ".log"}:
        return None
    if len(payload) > MAX_ATTACHMENT_BYTES:
        return "附件超过 5 MB，未自动分析。"
    try:
        if suffix == ".xlsx":
            return _analyze_xlsx(payload)
        if suffix == ".docx":
            return _analyze_docx(payload)
        text = payload.decode("utf-8-sig", errors="replace")
        if suffix == ".csv":
            rows = list(csv.reader(io.StringIO(text)))[:8]
            preview = "；".join(" | ".join(_cell(value) for value in row[:6]) for row in rows)
            return f"CSV 表格预览：{preview or '空文件'}"
        return "文本内容：" + (
            " ".join(text.split())[:MAX_ATTACHMENT_ANALYSIS_CHARS] or "空文件"
        )
    except Exception:
        return "附件格式异常，无法自动分析。"


def _analyze_xlsx(payload: bytes) -> str:
    if _unpacked_size(payload) > MAX_UNPACKED_BYTES:
        return "Excel 解压后超过 20 MB，未自动分析。"
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        parts = [f"Excel，共 {len(workbook.sheetnames)} 个工作表"]
        for name in workbook.sheetnames[:3]:
            sheet = workbook[name]
            rows = [
                [_cell(value) for value in row]
                for row in sheet.iter_rows(values_only=True)
                if any(value not in (None, "") for value in row)
            ]
            if not rows:
                parts.append(f"{name}：空表")
                continue
            header_index = max(
                range(min(20, len(rows))),
                key=lambda index: sum(bool(value) for value in rows[index]),
            )
            header = rows[header_index]
            columns = [
                index
                for index, value in enumerate(header)
                if value or any(index < len(row) and row[index] for row in rows[header_index + 1 :])
            ]
            data = [
                row
                for row in rows[header_index + 1 :]
                if sum(bool(row[index]) for index in columns if index < len(row)) >= min(2, len(columns))
            ]
            labels = [(header[index] or f"第{index + 1}列")[:24] for index in columns]
            summary = f"{name}：{len(data)} 条有效数据，{len(columns)} 个字段"
            if labels:
                summary += "；字段：" + "、".join(labels)
            stats = []
            samples = []
            for index, label in zip(columns, labels):
                values = [row[index] for row in data if index < len(row) and row[index]]
                counts = Counter(values)
                if values and len(counts) <= min(8, max(2, len(values) // 2)):
                    stats.append(
                        f"{label}="
                        + "、".join(f"{value[:24]} {count}" for value, count in counts.most_common(3))
                    )
                elif len(values) >= 3 and sum(map(len, values)) / len(values) >= 8:
                    samples.append(f"{label}样本=" + "｜".join(value[:60] for value in values[:2]))
            if stats:
                summary += "；统计：" + "；".join(stats[:5])
            if samples:
                summary += "；开放反馈：" + "；".join(samples[:2])
            parts.append(summary)
        return "；".join(parts)[:MAX_ATTACHMENT_ANALYSIS_CHARS]
    finally:
        workbook.close()


def _analyze_docx(payload: bytes) -> str:
    if _unpacked_size(payload) > MAX_UNPACKED_BYTES:
        return "Word 解压后超过 20 MB，未自动分析。"
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    text = " ".join(
        str(node.text or "").strip()
        for node in root.iter()
        if node.tag.endswith("}t") and str(node.text or "").strip()
    )
    return "Word 文档内容：" + (text[:MAX_ATTACHMENT_ANALYSIS_CHARS] or "空文档")


def _unpacked_size(payload: bytes) -> int:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return sum(item.file_size for item in archive.infolist())


def _cell(value: Any) -> str:
    return " ".join(str(value if value is not None else "").split())[:80]


def _login(host: str, port: int, username: str, password: str, timeout: int) -> poplib.POP3_SSL:
    client: poplib.POP3_SSL | None = None
    try:
        client = poplib.POP3_SSL(
            host,
            port,
            timeout=max(1, timeout),
            context=ssl.create_default_context(),
        )
        client.user(username)
        client.pass_(password)
        return client
    except poplib.error_proto as exc:
        if client:
            _quit(client)
        raise EmailAuthenticationError("邮箱账号或密码错误") from exc
    except (OSError, ssl.SSLError) as exc:
        if client:
            _quit(client)
        raise EmailConnectionError("暂时无法连接邮箱服务器") from exc


def _quit(client: poplib.POP3_SSL) -> None:
    try:
        client.quit()
    except Exception:  # noqa: BLE001
        try:
            client.close()
        except Exception:  # noqa: BLE001
            pass


def _uidl_entry(line: bytes) -> tuple[int, str]:
    number, uidl = line.decode("ascii", errors="strict").split(maxsplit=1)
    return int(number), uidl.strip()


def _message_date(message: Message) -> datetime | None:
    value = message.get("Date")
    if not value:
        return None
    try:
        parsed = parsedate_to_datetime(str(value))
    except (TypeError, ValueError):
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc).replace(tzinfo=None)


def _addresses(message: Message, header: str) -> list[dict[str, str]]:
    return [
        {"name": name[:500], "address": address.lower()[:320]}
        for name, address in getaddresses([str(message.get(header) or "")])
        if address
    ]


def _part_text(part: Message, payload: bytes) -> str:
    charset = part.get_content_charset() or "utf-8"
    try:
        return payload.decode(charset, errors="replace")
    except LookupError:
        return payload.decode("utf-8", errors="replace")
