from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.file_analysis import analyze_uploaded_file_async


def build_upload_ack(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return build_xlsx_ack(path)
    return (
        f"已收到文件：{path.name}\n"
        f"大小：{path.stat().st_size} 字节\n\n"
        "你可以继续问我“帮我分析这份文件”或提出具体问题。"
    )


def build_xlsx_ack(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    lines = [f"已收到 Excel：{path.name}", f"工作表数：{len(wb.sheetnames)}"]
    for sheet_name in wb.sheetnames[:3]:
        ws = wb[sheet_name]
        lines.append(f"- {sheet_name}：{ws.max_row} 行 x {ws.max_column} 列")
    lines.append("\n你可以继续问我：帮我分析这份表、汇总销售趋势、按地区排名等。")
    return "\n".join(lines)


async def analyze_file_for_question(path: Path, question: str) -> str:
    if path.suffix.lower() == ".xlsx" and is_sales_trend_question(question):
        return analyze_sales_trend_xlsx(path, question)
    return await analyze_uploaded_file_async(path)


def is_sales_trend_question(text: str) -> bool:
    return "销售" in text and any(marker in text for marker in ("趋势", "分析", "汇总", "统计"))


def analyze_sales_trend_xlsx(path: Path, question: str) -> str:
    rows = read_first_sheet_rows(path)
    if not rows:
        return f"未能读取到 Excel 数据：{path.name}"

    headers = [str(value or "").strip() for value in rows[0]]
    records = [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]
    if not records:
        return f"Excel 没有可分析的数据行：{path.name}"

    date_col = find_column(headers, ("日期", "date", "时间"))
    amount_col = find_column(headers, ("销售额", "金额", "销售金额", "收入", "amount"))
    qty_col = find_column(headers, ("销售数量", "数量", "销量", "qty"))
    region_col = find_column(headers, ("销售地区", "地区", "区域", "region"))
    category_col = find_column(headers, ("产品类别", "类别", "品类", "category"))
    seller_col = find_column(headers, ("销售人员", "业务员", "人员", "seller"))

    if not amount_col and not qty_col:
        return (
            f"已读取 {path.name}，但没有识别到销售额/销售数量列。\n\n"
            f"识别到的列：{', '.join(headers)}"
        )

    metrics = build_sales_metrics(records, date_col, amount_col, qty_col, region_col, category_col, seller_col)
    return format_sales_report(path.name, question, metrics, amount_col, qty_col, region_col, category_col, seller_col)


def read_first_sheet_rows(path: Path) -> list[tuple[Any, ...]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return list(ws.iter_rows(values_only=True))


def find_column(headers: list[str], candidates: tuple[str, ...]) -> str | None:
    lowered = [(header, header.lower()) for header in headers]
    for candidate in candidates:
        candidate_lower = candidate.lower()
        for header, lower in lowered:
            if candidate_lower in lower:
                return header
    return None


def to_number(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").replace("￥", "").strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def to_month(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m")
    if isinstance(value, date):
        return value.strftime("%Y-%m")
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m", "%Y/%m"):
        try:
            return datetime.strptime(text[:10] if "d" in fmt else text[:7], fmt).strftime("%Y-%m")
        except ValueError:
            continue
    return "未识别"


def build_sales_metrics(
    records: list[dict[str, Any]],
    date_col: str | None,
    amount_col: str | None,
    qty_col: str | None,
    region_col: str | None,
    category_col: str | None,
    seller_col: str | None,
) -> dict[str, Any]:
    total_amount = sum(to_number(row.get(amount_col)) for row in records) if amount_col else 0.0
    total_qty = sum(to_number(row.get(qty_col)) for row in records) if qty_col else 0.0
    monthly: dict[str, dict[str, float]] = defaultdict(lambda: {"amount": 0.0, "qty": 0.0})
    region: dict[str, float] = defaultdict(float)
    category: dict[str, float] = defaultdict(float)
    seller: dict[str, float] = defaultdict(float)

    for row in records:
        amount = to_number(row.get(amount_col)) if amount_col else 0.0
        qty = to_number(row.get(qty_col)) if qty_col else 0.0
        month = to_month(row.get(date_col)) if date_col else "未识别"
        monthly[month]["amount"] += amount
        monthly[month]["qty"] += qty
        if region_col:
            region[str(row.get(region_col) or "未填写")] += amount or qty
        if category_col:
            category[str(row.get(category_col) or "未填写")] += amount or qty
        if seller_col:
            seller[str(row.get(seller_col) or "未填写")] += amount or qty

    return {
        "row_count": len(records),
        "total_amount": total_amount,
        "total_qty": total_qty,
        "monthly": dict(sorted(monthly.items())),
        "region": top_items(region),
        "category": top_items(category),
        "seller": top_items(seller),
    }


def top_items(values: dict[str, float], limit: int = 8) -> list[tuple[str, float]]:
    return sorted(values.items(), key=lambda item: item[1], reverse=True)[:limit]


def format_sales_report(
    file_name: str,
    question: str,
    metrics: dict[str, Any],
    amount_col: str | None,
    qty_col: str | None,
    region_col: str | None,
    category_col: str | None,
    seller_col: str | None,
) -> str:
    lines = [
        f"基于上传文件 **{file_name}** 的销售趋势分析",
        "",
        "## 总览",
        f"- 数据行数：{metrics['row_count']}",
    ]
    if amount_col:
        lines.append(f"- 总销售额：{metrics['total_amount']:,.2f}")
    if qty_col:
        lines.append(f"- 总销售数量：{metrics['total_qty']:,.0f}")

    lines.extend(["", "## 月度趋势", "| 月份 | 销售额 | 销售数量 |", "|---|---:|---:|"])
    monthly = metrics["monthly"]
    for month, value in monthly.items():
        lines.append(f"| {month} | {value['amount']:,.2f} | {value['qty']:,.0f} |")

    append_ranking(lines, "地区排名", region_col, metrics["region"])
    append_ranking(lines, "产品类别排名", category_col, metrics["category"])
    append_ranking(lines, "销售人员排名", seller_col, metrics["seller"])

    months = [month for month in monthly if month != "未识别"]
    if len(months) >= 2 and amount_col:
        first = monthly[months[0]]["amount"]
        last = monthly[months[-1]]["amount"]
        change = ((last - first) / first * 100) if first else 0
        lines.extend(["", "## 结论", f"- 从 {months[0]} 到 {months[-1]}，销售额变化约 {change:.1f}%。"])
    lines.append("- 如果要进一步分析，可以继续问：按地区看趋势、按产品类别看趋势、找销售异常月份。")
    return "\n".join(lines)


def append_ranking(lines: list[str], title: str, column: str | None, rows: list[tuple[str, float]]) -> None:
    if not column or not rows:
        return
    lines.extend(["", f"## {title}", f"| {column} | 指标值 |", "|---|---:|"])
    for name, value in rows:
        lines.append(f"| {name} | {value:,.2f} |")
